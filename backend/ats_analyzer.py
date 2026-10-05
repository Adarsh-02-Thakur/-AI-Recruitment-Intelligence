"""Resume-only ATS analyzer. Fixed 100-point rubric: same resume -> same score, every lost point is explained."""
import io, math, re
import docx, pdfplumber
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from scoring import extract_skills, skill_gap

router = APIRouter(prefix="/api/ats")
STRONG = set("""built developed designed led created improved reduced increased launched implemented optimized automated managed delivered
achieved analyzed architected deployed engineered established generated mentored migrated negotiated streamlined trained resolved
collaborated drove spearheaded integrated tested maintained coordinated presented researched secured scaled supervised initiated""".split())
WEAK = ["responsible for", "worked on", "helped with", "helped", "duties included", "assisted with", "tasked with", "participated in"]
MON = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s+"
RANGE = re.compile(rf"(?:{MON})?(?:19|20)\d{{2}}\s*(?:-|–|—|to)\s*(?:(?:{MON})?(?:19|20)\d{{2}}|present|current|now)", re.I)
SECT = {"summary": r"^\s*(professional\s+)?(summary|objective|profile|about me)\b", "experience": r"^\s*(work\s+|professional\s+)?(experience|employment|work history|internships?)\b",
        "education": r"^\s*education\b", "skills": r"^\s*(technical\s+|key\s+|core\s+)?(skills|competencies)\b",
        "projects/certifications": r"^\s*(projects?|certifications?|achievements?|awards)\b"}


def read_file(name: str, data: bytes):
    n = name.lower(); meta = {"type": n.rsplit(".", 1)[-1], "pages": None, "tables": 0, "images": 0}
    if n.endswith(".pdf"):
        parts = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            meta["pages"] = len(pdf.pages)
            for p in pdf.pages:
                parts.append(p.extract_text() or ""); meta["tables"] += len(p.find_tables()); meta["images"] += len(p.images)
        text = "\n".join(parts)
    elif n.endswith(".docx"):
        d = docx.Document(io.BytesIO(data)); meta["tables"] = len(d.tables); meta["images"] = len(d.inline_shapes)
        text = "\n".join(p.text for p in d.paragraphs)
        for t in d.tables:
            for row in t.rows: text += "\n" + " | ".join(c.text for c in row.cells)
    else:
        text = data.decode("utf-8", errors="ignore")
    meta["scanned"] = meta["type"] == "pdf" and len(text.strip()) < 100
    return text, meta


def analyze(text: str, meta: dict | None = None, jd: str = "") -> dict:
    meta = meta or {}; low = text.lower()
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    clines = [l for l in lines if len(l.split()) >= 6]
    wc = len(re.findall(r"[A-Za-z0-9+#.']+", text))
    checks = []

    def add(cat, name, pts, mx, tip):
        pts = round(min(pts, mx), 1)
        checks.append({"cat": cat, "name": name, "pts": pts, "max": mx, "tip": "" if pts >= mx else tip,
                       "status": "pass" if pts >= mx else "warn" if pts >= mx * 0.5 else "fail"})

    # 1 FORMAT (20)
    pages = meta.get("pages") or max(1, math.ceil(wc / 500))
    add("Format & Parsing", "Text is machine-readable", 0 if (meta.get("scanned") or wc < 80) else 6, 6, "Little or no text could be read. Scanned or image-only resumes are invisible to ATS. Export a text-based PDF or a DOCX.")
    add("Format & Parsing", "Length in pages (1-2 ideal)", 4 if pages <= 2 else 2 if pages == 3 else 0, 4, f"Your resume is about {pages} pages. Keep it to 1-2 pages.")
    add("Format & Parsing", "No tables or columns", 4 if meta.get("tables", 0) == 0 else 1, 4, "Tables and multi-column layouts can scramble the reading order. Use a single-column layout.")
    add("Format & Parsing", "No images or graphics", 3 if meta.get("images", 0) == 0 else 1 if meta.get("images") == 1 else 0, 3, "Photos, logos and skill-bar graphics are ignored by ATS. Remove them.")
    odd = len(re.findall("[\u2600-\u27BF\U0001F300-\U0001FAFF\uE000-\uF8FF]", text))
    add("Format & Parsing", "No emoji or icon symbols", 3 if odd == 0 else 0, 3, "Emoji and icon fonts show up as garbage characters. Use plain text and simple bullets.")

    # 2 CONTACT (10)
    add("Contact Info", "Email address", 4 if re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text) else 0, 4, "Add a professional email address.")
    add("Contact Info", "Phone number", 3 if re.search(r"\+?\d[\d\s().-]{8,}\d", text) else 0, 3, "Add a phone number.")
    li, gh = "linkedin.com/" in low, bool(re.search(r"github\.com/|portfolio|behance\.net/|\.dev\b", low))
    add("Contact Info", "LinkedIn / portfolio link", (2 if li else 0) + (1 if gh else 0), 3, "Add your LinkedIn URL and a GitHub or portfolio link.")

    # 3 SECTIONS (15)
    for key, mx in (("summary", 3), ("experience", 4), ("education", 3), ("skills", 3), ("projects/certifications", 2)):
        add("Sections", f"'{key.title()}' section", mx if re.search(SECT[key], text, re.I | re.M) else 0, mx,
            f"Add a clearly labelled '{key.split('/')[0].title()}' heading. ATS looks for standard section names.")

    # 4 CONTENT (25)
    verbs = sum(1 for v in STRONG if re.search(rf"\b{v}\b", low))
    add("Content Quality", "Strong action verbs", min(verbs / 8, 1) * 8, 8, f"Only {verbs} distinct strong verbs found. Start bullets with verbs like Built, Led, Reduced, Automated.")
    clean = re.sub(r"(?:19|20)\d{2}|\+?\d[\d\s().-]{8,}\d", " ", low)
    metrics = len(re.findall(r"\d+(?:\.\d+)?\s*(?:%|\+|x\b|k\b|million|lakh|crore|users|clients|customers|ms\b|hours)", clean)) + len(re.findall(r"[$₹€£]\s*\d", clean))
    add("Content Quality", "Quantified achievements", min(metrics / 6, 1) * 8, 8, f"Only {metrics} measurable results found. Add numbers: %, time saved, users, revenue.")
    weak = sum(low.count(w) for w in WEAK)
    add("Content Quality", "No weak phrases", 4 if weak == 0 else 2 if weak <= 2 else 0, 4, f"Found {weak} weak phrases like 'responsible for' or 'worked on'. Replace them with action verbs.")
    pron = len(re.findall(r"\b(i|my|me|myself)\b", low))
    add("Content Quality", "No first-person pronouns", 2 if pron == 0 else 1 if pron <= 2 else 0, 2, "Remove 'I' and 'my'. Resumes use implied first person.")
    add("Content Quality", "Enough descriptive lines", 3 if len(clines) >= 10 else 2 if len(clines) >= 5 else 0, 3, "Add more bullet points describing what you did and the result.")

    # 5 SKILLS & KEYWORDS (15)
    skills = sorted(extract_skills(text)); jd_gap = None
    if jd.strip():
        jd_gap = skill_gap(text, jd)
        add("Skills & Keywords", "Job-description keyword match", jd_gap["coverage"] / 100 * 15, 15,
            "Missing keywords from the job: " + (", ".join(jd_gap["missing"][:10]) or "none") + ". Add them only if you truly have them.")
    else:
        add("Skills & Keywords", "Skills coverage (no job description given)", min(len(skills) / 10, 1) * 15, 15,
            f"Only {len(skills)} recognised skills found. List more tools, languages and frameworks you know.")

    # 6 READABILITY (10)
    add("Readability & Length", "Word count (350-800 ideal)", 5 if 350 <= wc <= 800 else 3 if 250 <= wc <= 1000 else 1 if wc else 0, 5, f"Your resume has {wc} words. Aim for 350-800.")
    avg = sum(len(l.split()) for l in clines) / len(clines) if clines else 0
    add("Readability & Length", "Concise bullet length", 3 if 8 <= avg <= 28 else 1, 3, "Keep bullets to roughly 1-2 lines (8-28 words).")
    caps = [w for w in re.findall(r"\b[A-Z]{3,}\b", text)]
    add("Readability & Length", "Limited ALL-CAPS text", 2 if len(caps) <= max(10, wc * 0.08) else 0, 2, "Too much ALL-CAPS text. Use it only for section headings.")

    # 7 DATES (5)
    ranges = RANGE.findall(text); styles = {"month" if re.search(MON, r, re.I) else "year" for r in ranges}
    add("Dates & Timeline", "Employment / education date ranges", 3 if len(ranges) >= 2 else 1 if ranges else 0, 3, "Add dates for each role or degree, like 'Jun 2023 - Present'.")
    add("Dates & Timeline", "Consistent date format", 2 if ranges and len(styles) == 1 else 0, 2, "Use one date style everywhere, such as 'Jun 2023' or '2023'.")

    cats = {}
    for c in checks:
        d = cats.setdefault(c["cat"], {"score": 0.0, "max": 0}); d["score"] = round(d["score"] + c["pts"], 1); d["max"] += c["max"]
    score = round(sum(c["pts"] for c in checks), 1)
    grade = "Excellent" if score >= 90 else "Good" if score >= 75 else "Fair" if score >= 60 else "Needs work"
    fixes = sorted([c for c in checks if c["max"] - c["pts"] > 0], key=lambda c: c["pts"] - c["max"])[:5]
    return {"score": score, "grade": grade, "categories": cats, "checks": checks, "skills": skills,
            "jd_missing": jd_gap["missing"] if jd_gap else None,
            "top_fixes": [{"name": f["name"], "lost": round(f["max"] - f["pts"], 1), "tip": f["tip"]} for f in fixes],
            "stats": {"words": wc, "pages": pages, "metrics": metrics, "action_verbs": verbs, "skills_found": len(skills)}}


@router.post("/analyze")
async def analyze_endpoint(file: UploadFile | None = File(None), text: str = Form(""), jd: str = Form("")):
    meta = {}
    if file and file.filename:
        data = await file.read()
        if len(data) > 8_000_000: raise HTTPException(413, "File too large (max 8 MB)")
        try: text, meta = read_file(file.filename, data)
        except Exception: raise HTTPException(422, "Could not read this file. Try PDF, DOCX or TXT.")
    if len(text.strip()) < 20 and not meta.get("scanned"): raise HTTPException(422, "Upload a resume or paste its text")
    return analyze(text, meta, jd)